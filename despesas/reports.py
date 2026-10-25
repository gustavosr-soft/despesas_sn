"""Relatórios (um dia, um mês ou vários meses) em Excel (.xlsx) e PDF."""
from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (
    HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from . import money

MIN_LINHAS = 15  # Excel: completa com linhas em branco

MESES_PT = [
    "", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]


def _br(d) -> str:
    return d.strftime("%d/%m/%Y")


def _intervalo_txt(r: dict) -> str:
    ini, fim = r.get("inicio"), r.get("fim")
    if ini is None:
        return f"{r['mes']:02d}/{r['ano']}"
    return _br(ini) if ini == fim else f"{_br(ini)} a {_br(fim)}"


def _mes_inteiro(r: dict) -> bool:
    return r.get("mes_inteiro", True)


def titulo(r: dict) -> str:
    if _mes_inteiro(r):
        return f"DESPESAS {r['mes']:02d}/{r['ano']}"
    return f"DESPESAS {_intervalo_txt(r)}"


def titulo_extenso(r: dict) -> str:
    if _mes_inteiro(r):
        return f"Despesas — {MESES_PT[r['mes']]} de {r['ano']}"
    return f"Despesas — {_intervalo_txt(r)}"


def nome_arquivo(r: dict, extensao: str) -> str:
    if _mes_inteiro(r):
        return f"despesas_{r['ano']}-{r['mes']:02d}.{extensao}"
    ini, fim = r["inicio"], r["fim"]
    return f"despesas_{ini.isoformat()}.{extensao}" if ini == fim else \
        f"despesas_{ini.isoformat()}_a_{fim.isoformat()}.{extensao}"


def _rotulo_mes(chave: str) -> str:
    ano, mes = chave.split("-")
    return f"{MESES_PT[int(mes)]} de {ano}"


def _texto_vazio(r: dict) -> str:
    return "Nenhum lançamento registrado neste mês." if _mes_inteiro(r) and not r.get("filtros_txt") \
        else "Nenhum lançamento encontrado neste período/filtro."


def _dia_txt(l: dict) -> str:
    return l.get("dia_txt") or f"{l['dia']:02d}"


def _cols_com_dados(r: dict) -> list:
    """Retorna apenas colunas que possuem ao menos um lançamento no mês."""
    usadas = {l["categoria_id"] for l in r["linhas"]}
    return [c for c in r["colunas"] if c["id"] in usadas]


def linhas_com_branco(r: dict) -> list:
    """Lançamentos + linhas em branco até MIN_LINHAS (None = linha em branco)."""
    linhas = list(r["linhas"])
    return linhas + [None] * max(0, MIN_LINHAS - len(linhas))


# ----------------------------------------------------------------------- Excel
def gerar_xlsx(r: dict, empresa: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{r['mes']:02d}-{r['ano']}" if _mes_inteiro(r) else "Período"

    # Filtra colunas — só as que têm lançamentos
    cols = _cols_com_dados(r)

    if not cols:
        # Mês sem lançamentos
        ws["A1"] = titulo(r)
        ws["A1"].font = Font(bold=True, size=16)
        ws["A3"] = _texto_vazio(r)
        ws["A3"].font = Font(size=12, italic=True)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    ncol = 2 + len(cols)
    fino = Side(style="thin", color="000000")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    cinza = PatternFill("solid", fgColor="E8E8E8")
    azul_cab = PatternFill("solid", fgColor="2B4C7E")
    branco = Font(bold=True, size=10, color="FFFFFF")
    fmt_valor = '#,##0.00;-#,##0.00;"—"'

    # Título e empresa
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A1"] = titulo(r)
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    sub = " — ".join(x for x in (empresa, r.get("filtros_txt", "")) if x)
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        ws["A2"] = sub
        ws["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
        ws["A2"].font = Font(size=10, italic=True)
        ws.row_dimensions[2].height = 28 if len(sub) > 90 else 15

    cab = 3
    nomes_cab = ["DIA", "DESCRIÇÃO"] + [c["nome"].upper() for c in cols]
    for i, nome in enumerate(nomes_cab, start=1):
        c = ws.cell(row=cab, column=i, value=nome)
        c.font = branco
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = azul_cab
        c.border = borda
    ws.row_dimensions[cab].height = 26

    linhas = linhas_com_branco(r)
    col_ids = [c_["id"] for c_ in cols]
    primeira = cab + 1
    zebra_claro = PatternFill("solid", fgColor="F5F7FA")
    for i, l in enumerate(linhas):
        linha = primeira + i
        for col in range(1, ncol + 1):
            cel = ws.cell(row=linha, column=col)
            cel.border = borda
            if i % 2 == 1:
                cel.fill = zebra_claro
        if l is None:
            continue
        c = ws.cell(row=linha, column=1, value=l["dia"] if not r.get("multi_mes") else _dia_txt(l))
        c.number_format = "00"
        c.alignment = Alignment(horizontal="center")
        d = ws.cell(row=linha, column=2, value=l["descricao"])
        if str(l["descricao"]).startswith("="):
            d.data_type = "s"
        d.alignment = Alignment(wrap_text=True, vertical="top")
        if l["categoria_id"] in col_ids:
            idx = col_ids.index(l["categoria_id"])
            v = ws.cell(row=linha, column=3 + idx, value=round(l["valor"] / 100, 2))
            v.number_format = fmt_valor
    ultima = primeira + len(linhas) - 1

    # Linha de totais
    tot = ultima + 1
    tot_fill = PatternFill("solid", fgColor="D6E4F0")
    ws.cell(row=tot, column=1).border = borda
    ws.cell(row=tot, column=1).fill = tot_fill
    t = ws.cell(row=tot, column=2, value="TOTAIS")
    t.font = Font(bold=True, size=12)
    t.alignment = Alignment(horizontal="center")
    t.border = borda
    t.fill = tot_fill
    for i in range(len(cols)):
        letra = get_column_letter(3 + i)
        c = ws.cell(row=tot, column=3 + i, value=f"=SUM({letra}{primeira}:{letra}{ultima})")
        c.number_format = fmt_valor
        c.font = Font(bold=True)
        c.border = borda
        c.fill = tot_fill

    # Total geral
    geral = tot + 2
    g = ws.cell(row=geral, column=2, value="TOTAL GERAL")
    g.font = Font(bold=True, size=13)
    g.alignment = Alignment(horizontal="right")
    primeiro_letra, ultima_letra = get_column_letter(3), get_column_letter(ncol)
    v = ws.cell(row=geral, column=3, value=f"=SUM({primeiro_letra}{tot}:{ultima_letra}{tot})")
    v.number_format = fmt_valor
    v.font = Font(bold=True, size=13)
    v.border = Border(bottom=fino)

    # Larguras
    ws.column_dimensions["A"].width = 8 if r.get("multi_mes") else 6
    ws.column_dimensions["B"].width = 38
    for i in range(len(cols)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 15

    # Impressão
    ws.freeze_panes = ws.cell(row=primeira, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{cab}:{cab}"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "Página &P de &N"

    if r.get("multi_mes") and r.get("por_mes"):  # segunda aba: total de cada mês
        wm = wb.create_sheet("Por mês")
        wm["A1"], wm["B1"] = "MÊS", "TOTAL (R$)"
        for c_ in (wm["A1"], wm["B1"]):
            c_.font = branco
            c_.fill = azul_cab
            c_.border = borda
            c_.alignment = Alignment(horizontal="center")
        for n, (chave, total) in enumerate(r["por_mes"], start=2):
            wm.cell(row=n, column=1, value=_rotulo_mes(chave)).border = borda
            v_ = wm.cell(row=n, column=2, value=round(total / 100, 2))
            v_.number_format = fmt_valor
            v_.border = borda
        fim_ = 1 + len(r["por_mes"])
        wm.cell(row=fim_ + 1, column=1, value="TOTAL GERAL").font = Font(bold=True)
        t_ = wm.cell(row=fim_ + 1, column=2, value=f"=SUM(B2:B{fim_})")
        t_.number_format = fmt_valor
        t_.font = Font(bold=True)
        wm.column_dimensions["A"].width = 24
        wm.column_dimensions["B"].width = 18

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------- PDF
COR_AZUL = colors.HexColor("#2B4C7E")
COR_AZUL_CLARO = colors.HexColor("#D6E4F0")
COR_CINZA_CLARO = colors.HexColor("#F5F7FA")
COR_CINZA = colors.HexColor("#888888")
COR_LINHA = colors.HexColor("#DDDDDD")

# Mesma paleta usada nos cards de resumo da tela (var(--cat-cor-N) em estilo.css)
PALETA_CATEGORIAS = [
    colors.HexColor("#3b82f6"), colors.HexColor("#ef4444"), colors.HexColor("#8b5cf6"),
    colors.HexColor("#f59e0b"), colors.HexColor("#10b981"), colors.HexColor("#ec4899"),
    colors.HexColor("#6b7280"), colors.HexColor("#14b8a6"), colors.HexColor("#f97316"),
    colors.HexColor("#6366f1"),
]

ESTILOS_PDF = {
    "lista": "Lista cronológica",
    "categoria": "Agrupado por categoria",
    "executivo_lista": "Resumo executivo + lista",
    "executivo_grafico": "Resumo executivo + gráfico",
}
ESTILO_PADRAO = "lista"


def _cor_categoria(idx: int):
    return PALETA_CATEGORIAS[idx % len(PALETA_CATEGORIAS)]


def _canvas_numerado(rodape_esq: str):
    class NumeradoCanvas(rl_canvas.Canvas):
        def __init__(self, *a, **kw):
            super().__init__(*a, **kw)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self.setFont("Helvetica", 7)
                self.setFillColor(COR_CINZA)
                largura, _ = self._pagesize
                self.drawString(15 * mm, 10 * mm, rodape_esq)
                self.drawRightString(largura - 15 * mm, 10 * mm,
                                     f"Página {self._pageNumber} de {total}")
                self.setStrokeColor(COR_LINHA)
                self.setLineWidth(0.3)
                self.line(15 * mm, 13 * mm, largura - 15 * mm, 13 * mm)
                super().showPage()
            super().save()

    return NumeradoCanvas


class _Estilos:
    """Agrupa os ParagraphStyles usados nos relatórios PDF."""

    def __init__(self):
        self.titulo = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=16,
                                     alignment=TA_CENTER, leading=20, textColor=COR_AZUL)
        self.sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=10,
                                  alignment=TA_CENTER, leading=13, textColor=COR_CINZA)
        self.cab = ParagraphStyle("cab", fontName="Helvetica-Bold", fontSize=8.5,
                                  alignment=TA_CENTER, leading=11, textColor=colors.white)
        self.txt = ParagraphStyle("txt", fontName="Helvetica", fontSize=9,
                                  alignment=TA_LEFT, leading=12)
        self.txt_compacto = ParagraphStyle("txtc", fontName="Helvetica", fontSize=8,
                                           alignment=TA_LEFT, leading=10.5)
        self.num = ParagraphStyle("num", fontName="Helvetica", fontSize=9,
                                  alignment=TA_RIGHT, leading=12)
        self.num_compacto = ParagraphStyle("numc", fontName="Helvetica", fontSize=8,
                                           alignment=TA_RIGHT, leading=10.5)
        self.dia = ParagraphStyle("dia", fontName="Helvetica", fontSize=9,
                                  alignment=TA_CENTER, leading=12)
        self.dia_compacto = ParagraphStyle("diac", fontName="Helvetica", fontSize=8,
                                           alignment=TA_CENTER, leading=10.5)
        self.tot_label = ParagraphStyle("tot_l", fontName="Helvetica-Bold", fontSize=10,
                                        alignment=TA_LEFT, leading=13)
        self.tot_val = ParagraphStyle("tot_v", fontName="Helvetica-Bold", fontSize=10,
                                      alignment=TA_RIGHT, leading=13)
        self.geral_label = ParagraphStyle("gl", fontName="Helvetica-Bold", fontSize=13,
                                          alignment=TA_LEFT, leading=16, textColor=COR_AZUL)
        self.geral_val = ParagraphStyle("gv", fontName="Helvetica-Bold", fontSize=13,
                                        alignment=TA_RIGHT, leading=16, textColor=COR_AZUL)
        self.sec = ParagraphStyle("sec", fontName="Helvetica-Bold", fontSize=11,
                                  alignment=TA_LEFT, leading=14, textColor=COR_AZUL,
                                  spaceBefore=8, spaceAfter=4)
        self.vazio = ParagraphStyle("vazio", fontName="Helvetica-Oblique", fontSize=11,
                                    alignment=TA_CENTER, leading=14, textColor=COR_CINZA)
        self.card_nome = ParagraphStyle("cn", fontName="Helvetica-Bold", fontSize=7.5,
                                        alignment=TA_LEFT, leading=9.5, textColor=COR_CINZA)
        self.card_valor = ParagraphStyle("cv", fontName="Helvetica-Bold", fontSize=13,
                                         alignment=TA_LEFT, leading=16)
        self.cat_titulo = ParagraphStyle("ct", fontName="Helvetica-Bold", fontSize=10.5,
                                         alignment=TA_LEFT, leading=13)
        self.subtotal = ParagraphStyle("subt", fontName="Helvetica-Bold", fontSize=9,
                                       alignment=TA_RIGHT, leading=12, textColor=COR_AZUL)


def _cabecalho(r: dict, empresa: str, est: "_Estilos") -> list:
    story = [Paragraph(titulo_extenso(r), est.titulo)]
    if empresa:
        story.append(Paragraph(escape(empresa), est.sub))
    if r.get("filtros_txt"):
        story.append(Paragraph(escape("Filtros: " + r["filtros_txt"]), est.sub))
    story.append(Spacer(1, 3 * mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=COR_AZUL,
                             spaceBefore=2, spaceAfter=6))
    return story


def _pagina_vazia(r: dict, empresa: str, usuario: str, gerado_em: str, est: "_Estilos") -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=18 * mm,
        title=titulo(r), author="Sistema de Despesas",
    )
    story = _cabecalho(r, empresa, est)
    story.append(Spacer(1, 20 * mm))
    story.append(Paragraph(_texto_vazio(r), est.vazio))
    rodape = f"Gerado em {gerado_em}" + (f" por {usuario}" if usuario else "")
    doc.build(story, canvasmaker=_canvas_numerado(rodape))
    return buf.getvalue()


def _tabela_lancamentos(r: dict, doc, est: "_Estilos", com_categoria: bool = True,
                        compacta: bool = False) -> Table:
    """Tabela cronológica de lançamentos (já vêm ordenados por dia e, no mesmo
    dia, do maior para o menor valor). Se com_categoria, inclui a coluna Categoria."""
    cat_nomes = {c["id"]: c["nome"] for c in r["colunas"]}
    txt = est.txt_compacto if compacta else est.txt
    num = est.num_compacto if compacta else est.num
    dia = est.dia_compacto if compacta else est.dia

    if com_categoria:
        cabecalho = [Paragraph("DIA", est.cab), Paragraph("DESCRIÇÃO", est.cab),
                     Paragraph("CATEGORIA", est.cab), Paragraph("VALOR (R$)", est.cab)]
    else:
        cabecalho = [Paragraph("DIA", est.cab), Paragraph("DESCRIÇÃO", est.cab),
                     Paragraph("VALOR (R$)", est.cab)]

    dados = [cabecalho]
    for l in r["linhas"]:
        linha = [Paragraph(_dia_txt(l), dia), Paragraph(escape(l["descricao"]), txt)]
        if com_categoria:
            linha.append(Paragraph(escape(cat_nomes.get(l["categoria_id"], "?")), txt))
        linha.append(Paragraph(money.formatar(l["valor"]), num))
        dados.append(linha)

    w_dia = 15 * mm if r.get("multi_mes") else 12 * mm
    w_val = 26 * mm
    if com_categoria:
        w_cat = 32 * mm
        larguras = [w_dia, doc.width - w_dia - w_val - w_cat, w_cat, w_val]
    else:
        larguras = [w_dia, doc.width - w_dia - w_val, w_val]

    tabela = Table(dados, colWidths=larguras, repeatRows=1)
    n_linhas = len(dados)
    pad = 3 if compacta else 4
    estilos = [
        ("BACKGROUND", (0, 0), (-1, 0), COR_AZUL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), pad),
        ("BOTTOMPADDING", (0, 0), (-1, -1), pad),
        ("LEFTPADDING", (0, 0), (-1, -1), pad + 1),
        ("RIGHTPADDING", (0, 0), (-1, -1), pad + 1),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, COR_AZUL),
    ]
    for i in range(1, n_linhas):
        if i % 2 == 0:
            estilos.append(("BACKGROUND", (0, i), (-1, i), COR_CINZA_CLARO))
        estilos.append(("LINEBELOW", (0, i), (-1, i), 0.3, COR_LINHA))
    tabela.setStyle(TableStyle(estilos))
    return tabela


def _tabela_por_mes(r: dict, doc, est: "_Estilos") -> Table:
    dados = [[Paragraph("MÊS", est.cab), Paragraph("TOTAL (R$)", est.cab)]]
    for chave, total in r["por_mes"]:
        dados.append([Paragraph(_rotulo_mes(chave), est.txt), Paragraph(money.formatar(total), est.num)])
    w_val = 34 * mm
    tab = Table(dados, colWidths=[doc.width - w_val, w_val], repeatRows=1)
    estilos = [("BACKGROUND", (0, 0), (-1, 0), COR_AZUL), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
               ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    for i in range(1, len(dados)):
        if i % 2 == 0:
            estilos.append(("BACKGROUND", (0, i), (-1, i), COR_CINZA_CLARO))
        estilos.append(("LINEBELOW", (0, i), (-1, i), 0.3, COR_LINHA))
    tab.setStyle(TableStyle(estilos))
    return tab


def _tabela_resumo_categoria(r: dict, doc, est: "_Estilos", cols_com_dados: list) -> Table:
    dados = [[Paragraph(escape(c["nome"]), est.tot_label),
              Paragraph(money.formatar(r["totais"][c["id"]]), est.tot_val)]
             for c in cols_com_dados]
    w1, w2 = doc.width * 0.6, doc.width * 0.4
    tab = Table(dados, colWidths=[w1, w2])
    estilos = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]
    for i in range(len(dados)):
        if i % 2 == 0:
            estilos.append(("BACKGROUND", (0, i), (-1, i), COR_CINZA_CLARO))
        estilos.append(("LINEBELOW", (0, i), (-1, i), 0.3, COR_LINHA))
    tab.setStyle(TableStyle(estilos))
    return tab


def _barra_total_geral(r: dict, doc, est: "_Estilos") -> Table:
    w1, w2 = doc.width * 0.6, doc.width * 0.4
    tab = Table(
        [[Paragraph("TOTAL GERAL", est.geral_label),
          Paragraph(money.formatar(r["total_geral"]), est.geral_val)]],
        colWidths=[w1, w2],
    )
    tab.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, -1), COR_AZUL_CLARO),
    ]))
    return tab


def _card_categoria(nome: str, valor_centavos: int, cor, largura: float, est: "_Estilos") -> Table:
    t = Table(
        [[Paragraph(nome.upper(), est.card_nome)],
         [Paragraph(money.formatar(valor_centavos), est.card_valor)]],
        colWidths=[largura],
    )
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), COR_CINZA_CLARO),
        ("LINEBEFORE", (0, 0), (0, -1), 3, cor),
        ("TOPPADDING", (0, 0), (-1, 0), 7), ("BOTTOMPADDING", (0, 0), (-1, 0), 1),
        ("TOPPADDING", (0, 1), (-1, 1), 1), ("BOTTOMPADDING", (0, 1), (-1, 1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
    ]))
    return t


def _grade_cards(r: dict, doc, est: "_Estilos", cols_com_dados: list, por_linha: int = 3) -> Table:
    gap = 3 * mm
    largura_card = (doc.width - gap * (por_linha - 1)) / por_linha
    linhas_tabela, linha_atual = [], []
    for i, c in enumerate(cols_com_dados):
        linha_atual.append(_card_categoria(c["nome"], r["totais"][c["id"]], _cor_categoria(i), largura_card, est))
        if len(linha_atual) == por_linha:
            linhas_tabela.append(linha_atual)
            linha_atual = []
    if linha_atual:
        linha_atual.extend([""] * (por_linha - len(linha_atual)))
        linhas_tabela.append(linha_atual)

    grade = Table(linhas_tabela, colWidths=[largura_card] * por_linha)
    grade.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), gap),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), gap),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return grade


def _grafico_categorias(r: dict, doc, cols_com_dados: list, altura: float = 65 * mm):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing

    d = Drawing(doc.width, altura)
    chart = HorizontalBarChart()
    chart.x = 95
    chart.y = 10
    chart.height = altura - 20
    chart.width = doc.width - 115
    chart.data = [[r["totais"][c["id"]] / 100 for c in cols_com_dados]]
    chart.categoryAxis.categoryNames = [c["nome"] for c in cols_com_dados]
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 8
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 7
    chart.barWidth = 10
    chart.groupSpacing = 6
    chart.bars.strokeColor = None
    for i in range(len(cols_com_dados)):
        chart.bars[(0, i)].fillColor = _cor_categoria(i)
    d.add(chart)
    return d


def _corpo_lista(r: dict, doc, est: "_Estilos") -> list:
    """Cronológico: Dia | Descrição | Categoria | Valor, com resumo por categoria no final."""
    cols = _cols_com_dados(r)
    story = [Paragraph("Lançamentos", est.sec),
             _tabela_lancamentos(r, doc, est, com_categoria=True),
             Spacer(1, 6 * mm),
             Paragraph("Resumo por Categoria", est.sec),
             _tabela_resumo_categoria(r, doc, est, cols),
             Spacer(1, 4 * mm),
             HRFlowable(width="100%", thickness=1, color=COR_AZUL, spaceBefore=2, spaceAfter=4),
             _barra_total_geral(r, doc, est)]
    return story


def _corpo_categoria(r: dict, doc, est: "_Estilos") -> list:
    """Agrupado por categoria: seção por categoria com seus lançamentos e subtotal."""
    cols = _cols_com_dados(r)
    story = []
    for i, c in enumerate(cols):
        cor = _cor_categoria(i)
        linhas_cat = [l for l in r["linhas"] if l["categoria_id"] == c["id"]]

        titulo_tab = Table([[Paragraph(c["nome"], est.cat_titulo)]], colWidths=[doc.width])
        titulo_tab.setStyle(TableStyle([
            ("LINEBEFORE", (0, 0), (0, 0), 4, cor),
            ("BACKGROUND", (0, 0), (0, 0), COR_CINZA_CLARO),
            ("LEFTPADDING", (0, 0), (0, 0), 8),
            ("TOPPADDING", (0, 0), (0, 0), 4), ("BOTTOMPADDING", (0, 0), (0, 0), 4),
        ]))
        story.append(titulo_tab)
        story.append(Spacer(1, 1.5 * mm))

        dados = [[Paragraph("DIA", est.cab), Paragraph("DESCRIÇÃO", est.cab), Paragraph("VALOR (R$)", est.cab)]]
        for l in linhas_cat:
            dados.append([Paragraph(_dia_txt(l), est.dia), Paragraph(escape(l["descricao"]), est.txt),
                          Paragraph(money.formatar(l["valor"]), est.num)])
        w_dia, w_val = (17 if r.get("multi_mes") else 14) * mm, 30 * mm
        tab = Table(dados, colWidths=[w_dia, doc.width - w_dia - w_val, w_val], repeatRows=1)
        estilos = [
            ("BACKGROUND", (0, 0), (-1, 0), cor), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("LINEBELOW", (0, 0), (-1, 0), 0.6, cor),
        ]
        for j in range(1, len(dados)):
            if j % 2 == 0:
                estilos.append(("BACKGROUND", (0, j), (-1, j), COR_CINZA_CLARO))
            estilos.append(("LINEBELOW", (0, j), (-1, j), 0.3, COR_LINHA))
        tab.setStyle(TableStyle(estilos))
        story.append(tab)

        subtotal = Table(
            [[Paragraph(f"Subtotal {c['nome']}: {money.formatar(r['totais'][c['id']])}", est.subtotal)]],
            colWidths=[doc.width],
        )
        subtotal.setStyle(TableStyle([
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("LINEABOVE", (0, 0), (-1, -1), 0.4, COR_LINHA),
        ]))
        story.append(subtotal)
        story.append(Spacer(1, 5 * mm))

    story.append(HRFlowable(width="100%", thickness=1, color=COR_AZUL, spaceBefore=2, spaceAfter=4))
    story.append(_barra_total_geral(r, doc, est))
    return story


def _corpo_executivo_lista(r: dict, doc, est: "_Estilos") -> list:
    """Cards de totais por categoria + total geral em destaque, seguidos da lista compacta."""
    cols = _cols_com_dados(r)
    return [Paragraph("Resumo do Mês" if _mes_inteiro(r) else "Resumo do Período", est.sec),
            _grade_cards(r, doc, est, cols),
            Spacer(1, 3 * mm),
            _barra_total_geral(r, doc, est),
            Spacer(1, 7 * mm),
            Paragraph("Lançamentos", est.sec),
            _tabela_lancamentos(r, doc, est, com_categoria=True, compacta=True)]


def _corpo_executivo_grafico(r: dict, doc, est: "_Estilos") -> list:
    """Gráfico de barras por categoria + total geral em destaque, seguidos da lista compacta."""
    cols = _cols_com_dados(r)
    return [Paragraph("Resumo do Mês" if _mes_inteiro(r) else "Resumo do Período", est.sec),
            _grafico_categorias(r, doc, cols),
            Spacer(1, 3 * mm),
            _barra_total_geral(r, doc, est),
            Spacer(1, 7 * mm),
            Paragraph("Lançamentos", est.sec),
            _tabela_lancamentos(r, doc, est, com_categoria=True, compacta=True)]


_CORPOS = {
    "lista": _corpo_lista,
    "categoria": _corpo_categoria,
    "executivo_lista": _corpo_executivo_lista,
    "executivo_grafico": _corpo_executivo_grafico,
}


def gerar_pdf(r: dict, empresa: str = "", usuario: str = "", gerado_em: str = "",
             estilo: str = ESTILO_PADRAO) -> bytes:
    """Gera o relatório mensal em PDF no estilo escolhido (ver ESTILOS_PDF):
    - lista: cronológico (Dia|Descrição|Categoria|Valor) + resumo por categoria
    - categoria: agrupado por categoria, com subtotal de cada uma
    - executivo_lista: cards de totais por categoria + lista cronológica compacta
    - executivo_grafico: gráfico de barras por categoria + lista cronológica compacta
    Quando o mês não tem lançamentos, qualquer estilo gera uma página informando isso.
    """
    est = _Estilos()
    if not r["linhas"]:
        return _pagina_vazia(r, empresa, usuario, gerado_em, est)

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=18 * mm,
        title=titulo(r), author="Sistema de Despesas",
    )
    story = _cabecalho(r, empresa, est)
    if r.get("multi_mes") and len(r.get("por_mes", [])) > 1:
        story += [Paragraph("Total por mês", est.sec), _tabela_por_mes(r, doc, est), Spacer(1, 6 * mm)]
    story.extend(_CORPOS.get(estilo, _corpo_lista)(r, doc, est))

    rodape = f"Gerado em {gerado_em}" + (f" por {usuario}" if usuario else "")
    doc.build(story, canvasmaker=_canvas_numerado(rodape))
    return buf.getvalue()
